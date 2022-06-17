import openpyxl
import xlsxwriter
import io
from django.db import IntegrityError
from kronos.exceptions import AppLogicError

from client.models import Store
from client.service import client_service
from manager.models import City


def import_store_by_xlsx_sheet(client_id, xlsx_sheet)-> bool:
    store_data = get_store_data_from_xlsx_sheet(xlsx_sheet)
    client = client_service.find_client_by_id(client_id)

    city_list = [store['city'] for store in store_data if store['city'] is not None]
    state_list = [store['state'] for store in store_data if store['state'] is not None]

    city_exists_list = City.objects.filter(name__in = city_list, state__in = state_list).values_list('name', flat=True)
    unique_city_list = set(city_list)
    city_not_exists_list = unique_city_list.difference(set(city_exists_list))

    if city_not_exists_list:
        raise AppLogicError('{} cities are not found'.format(', '.join(city_not_exists_list)))

    code_list = [store['code'] for store in store_data if store['code'] is not None]
    code_exists_list = Store.objects.filter(client = client_id, code__in = code_list).values_list('code', flat=True)

    if code_exists_list:
        raise AppLogicError('{} code with this client already exists'.format(', '.join(code_exists_list)))

    store_list = []
    for store in store_data:
        city_obj = City.objects.filter(name = store['city'], state = store['state']).first()
        if not city_obj:
            raise AppLogicError('\"{}\" city is not found'.format(store['city']))
        obj_dict = {
            'code': store['code'] if store['code'] else '',
            'name': store['name'],
            'address': store['address'],
            'city': city_obj,
            'pincode': store['pincode'] if store['pincode'] else '',
            'phone': store['phone'] if store['phone'] else '',
            'map_location_link': store['map_location_link'] if store['map_location_link'] else '',
            'type': store['store_type'] if store['store_type'] else '',
            'priority': store['priority'] if store['priority'] else '',
            'client': client,
        }
        store_list.append(
            Store(**obj_dict)
        )
    try:
        Store.objects.bulk_create(store_list)
    except IntegrityError as e:
        raise AppLogicError("Client with code combination already exists")
    return True


def get_store_data_from_xlsx_sheet(xlsx_sheet):
    wb = openpyxl.load_workbook(xlsx_sheet, data_only=True)

    worksheet = wb.active
    row_count = worksheet.max_row
    column_count = worksheet.max_column

    headers = ['code', 'name', 'address', 'city', 'pincode', 'state', 'phone', 'map_location_link', 'store_type', 'priority']
    if not column_count == len(headers):
        raise AppLogicError('Invalid sheet data format')

    if row_count > 100:
        raise AppLogicError('Store count must be less than 100')

    for ind, head in enumerate(next(worksheet.iter_rows(min_row=1, max_row=1))):
        if not head.value == headers[ind]:
            raise AppLogicError('Invalid sheet header format')

    excel_data = []
    for row in worksheet.iter_rows(min_row=2):
        row_data = {}
        for cell in row:
            col_name = headers[cell.column - 1]
            row_data[col_name]= str(cell.value).strip() if cell.value else None
        if row_data:
            excel_data.append(row_data)

    return excel_data


def find_sample_xlsx_for_store_insert():
    output = io.BytesIO()
    workbook = xlsxwriter.Workbook(output, {'in_memory': True})
    worksheet = workbook.add_worksheet()

    sample = (
        ['code', 'name', 'address', 'city', 'pincode', 'state', 'phone', 'map_location_link', 'store_type', 'priority'],
        ['Store001', 'Sample store', 'Nagpur', 'Nagpur', '4400001', 'IN-MH', '1234567890', 'google map link', 'Example', '1'],
    )

    row = 0
    col = 0

    # Iterate over the data and write it out row by row.
    for item in (sample):
        col = 0
        for i in item:
            worksheet.write(row, col, i)
            col += 1
        row += 1

    workbook.close()
    output.seek(0)
    name = "Sample store insert report.xlsx"
    return output, name